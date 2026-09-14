#!/usr/bin/env python3
"""Extract PPTX content for agent-assisted redesign, not pixel-perfect conversion."""
from __future__ import annotations
import argparse
import hashlib
import json
from pathlib import Path
import sys


def extract(source: Path, destination: Path) -> dict:
    if not source.is_file(): raise FileNotFoundError(source)
    if source.suffix.lower()!='.pptx': raise ValueError('Only .pptx is supported; convert legacy .ppt locally first.')
    if destination.exists() and any(destination.iterdir()): raise FileExistsError('Extraction directory must be new or empty.')
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError as exc:
        raise RuntimeError('PPTX extraction needs: python -m pip install python-pptx') from exc
    deck=Presentation(source)
    destination.mkdir(parents=True,exist_ok=True);(destination/'images').mkdir(exist_ok=True)
    result={'source':source.name,'sha256':hashlib.sha256(source.read_bytes()).hexdigest(),
            'dimensions_emu':{'width':deck.slide_width,'height':deck.slide_height},
            'warning':'Content extraction only. Verify charts, equations, SmartArt, animation, embedded objects, image crops and layout against original slides.', 'slides':[]}
    for number,slide in enumerate(deck.slides,1):
        record={'number':number,'title':slide.shapes.title.text if slide.shapes.title else '', 'elements':[], 'notes':'', 'warnings':[]}
        if slide.has_notes_slide:
            frame=slide.notes_slide.notes_text_frame
            if frame: record['notes']=frame.text
        def visit(shapes):
            for shape in shapes:
                item={'name':shape.name,'bounds_emu':{'left':shape.left,'top':shape.top,'width':shape.width,'height':shape.height}}
                if shape.shape_type==MSO_SHAPE_TYPE.GROUP:
                    visit(shape.shapes);continue
                if shape.has_text_frame:
                    item.update(type='text',text=shape.text_frame.text)
                elif shape.has_table:
                    item.update(type='table',rows=[[c.text for c in row.cells] for row in shape.table.rows])
                elif shape.shape_type==MSO_SHAPE_TYPE.PICTURE:
                    image=shape.image
                    safe_ext=''.join(c for c in image.ext.lower() if c.isalnum()) or 'bin'
                    name=f'slide-{number:03}-image-{len(record["elements"])+1:03}.{safe_ext}'
                    (destination/'images'/name).write_bytes(image.blob)
                    item.update(type='image',path='images/'+name,
                                crop={'left':shape.crop_left,'top':shape.crop_top,'right':shape.crop_right,'bottom':shape.crop_bottom})
                elif shape.has_chart:
                    item.update(type='chart',series=[])
                    try:
                        item['series']=[{'name':s.name,'values':list(s.values)} for s in shape.chart.series]
                    except Exception:
                        record['warnings'].append(f'Chart data could not be extracted: {shape.name}')
                    record['warnings'].append(f'Review chart labels, axes and category mapping against source: {shape.name}')
                else:
                    item.update(type='unhandled',shape_type=str(shape.shape_type))
                    record['warnings'].append(f'Visual/manual review required: {shape.name}')
                record['elements'].append(item)
        visit(slide.shapes)
        result['slides'].append(record)
    (destination/'content.json').write_text(json.dumps(result,ensure_ascii=False,indent=2,default=str)+'\n',encoding='utf-8')
    return result


def main() -> int:
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('input',type=Path);parser.add_argument('output_dir',type=Path)
    args=parser.parse_args()
    try:
        result=extract(args.input,args.output_dir)
        print(f'Extracted {len(result["slides"])} slides to {args.output_dir}/content.json. Visual review is required.')
        return 0
    except Exception as exc:
        print(f'Error: {exc}',file=sys.stderr);return 1

if __name__=='__main__': raise SystemExit(main())
